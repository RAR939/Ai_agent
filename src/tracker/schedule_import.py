"""Роль 4 (задача 4.2): импорт расписания из .ics или Excel в конкретные занятия по датам.

В базу кладутся именно занятия (дата и время), а не шаблон «по вторникам».
Повторная загрузка заменяет расписание целиком. Ошибки собираются списком.
"""
import re
from datetime import date, datetime, time, timedelta
from pathlib import Path

from src import config
from src.schemas import Lesson, ScheduleImport
from src.tracker import common, db

REQUIRED_COLUMNS = ["дисциплина", "тип", "преподаватель", "аудитория",
                    "день_недели", "начало", "конец", "неделя"]

_DAYS = {"пн": 0, "понедельник": 0, "вт": 1, "вторник": 1, "ср": 2, "среда": 2,
         "чт": 3, "четверг": 3, "пт": 4, "пятница": 4, "сб": 5, "суббота": 5,
         "вс": 6, "воскресенье": 6}
_WEEKS = {"чет": "even", "чёт": "even", "четная": "even", "чётная": "even",
          "нечет": "odd", "нечёт": "odd", "нечетная": "odd", "нечётная": "odd",
          "каждая": "all", "каждую": "all", "все": "all", "": "all"}


# ---------- общее ----------
def _semester() -> tuple[date, date, str, set[date]]:
    sch = common.cfg()["schedule"]
    start = date.fromisoformat(str(sch["semester_start"]))
    end = date.fromisoformat(str(sch["semester_end"]))
    parity = str(sch["first_week_parity"])
    holidays = {date.fromisoformat(str(h)) for h in (sch["holidays"] or [])}
    return start, end, parity, holidays


def week_parity(day: date, semester_start: date, first_week_parity: str = "odd") -> str:
    """'odd' (нечёт) или 'even' (чёт). Отсчёт от начала семестра, а не от 1 января."""
    index = ((day - semester_start).days // 7) % 2           # 0 для первой недели
    first_odd = first_week_parity == "odd"
    return "odd" if (index == 0) == first_odd else "even"


def _lesson_type(text: str) -> str:
    t = text.lower().replace("ё", "е")
    if "лаб" in t:
        return "lab"
    if "лек" in t:
        return "lecture"
    if "сем" in t or "практ" in t or t.startswith("пр"):
        return "seminar"
    return "other"


def _clean(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value != value:           # NaN из pandas
        return None
    text = str(value).strip()
    return text or None


def _check(lessons: list[Lesson], errors: list[str]) -> None:
    """Общая проверка перед записью: начало раньше конца."""
    for les in lessons:
        if common.parse_dt(les.start) >= common.parse_dt(les.end):
            errors.append(f"Занятие {les.discipline} {les.start}: начало не раньше конца")


# ---------- Excel ----------
def _parse_time(value, row: int, column: str, errors: list[str]) -> time | None:
    """Время из ячейки Excel: 'ЧЧ:ММ', time, datetime или доля суток."""
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, float) and 0 <= value < 1:             # доля суток из Excel
        minutes = round(value * 24 * 60)
        return time(minutes // 60 % 24, minutes % 60)
    text = _clean(value)
    if text is None:
        errors.append(f"Строка {row}: не указано время в колонке «{column}»")
        return None
    m = re.fullmatch(r"(\d{1,2})[:.](\d{2})(?::\d{2})?", text)
    if m and int(m.group(1)) < 24 and int(m.group(2)) < 60:
        return time(int(m.group(1)), int(m.group(2)))
    errors.append(f"Строка {row}: не удалось разобрать время в колонке «{column}»: {text!r}")
    return None


def _expand_excel(path: str, errors: list[str]) -> list[Lesson]:
    import pandas as pd

    df = pd.read_excel(path, dtype=object)
    df.columns = [str(c).strip().lower().replace(" ", "_") for c in df.columns]
    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        errors.append("В файле нет обязательных колонок: " + ", ".join(missing))
        return []
    start, end, first_parity, holidays = _semester()
    lessons: list[Lesson] = []
    for idx, rec in enumerate(df.to_dict("records")):
        row = idx + 2                                           # строка 1 в Excel это заголовок
        if all(_clean(v) is None for v in rec.values()):
            continue
        raw_disc = _clean(rec["дисциплина"])
        disc = config.find_discipline(raw_disc) if raw_disc else None
        n_before = len(errors)
        if raw_disc is None:
            errors.append(f"Строка {row}: не указана дисциплина")
        elif disc is None:
            errors.append(f'Строка {row}: не найдена дисциплина "{raw_disc}", '
                          f"добавьте алиас в disciplines.yaml")
        day_raw = (_clean(rec["день_недели"]) or "").lower().replace("ё", "е")
        weekday = _DAYS.get(day_raw)
        if weekday is None:
            errors.append(f"Строка {row}: непонятный день недели {day_raw!r}")
        week_raw = (_clean(rec["неделя"]) or "").lower().replace("ё", "е")
        week = _WEEKS.get(week_raw, _WEEKS.get(week_raw.replace("ё", "е")))
        if week is None:
            errors.append(f"Строка {row}: колонка «неделя» должна быть чёт, нечёт или каждая, "
                          f"а не {week_raw!r}")
        t_start = _parse_time(rec["начало"], row, "начало", errors)
        t_end = _parse_time(rec["конец"], row, "конец", errors)
        if len(errors) > n_before:
            continue
        if t_start >= t_end:
            errors.append(f"Строка {row}: начало ({t_start:%H:%M}) не раньше конца ({t_end:%H:%M})")
            continue
        lesson_type = _lesson_type(_clean(rec["тип"]) or "")
        teacher, room = _clean(rec["преподаватель"]), _clean(rec["аудитория"])
        day = start
        while day <= end:
            if (day.weekday() == weekday and day not in holidays
                    and (week == "all" or week_parity(day, start, first_parity) == week)):
                lessons.append(Lesson(
                    discipline=disc, lesson_type=lesson_type, teacher=teacher, room=room,
                    start=common.to_iso(datetime.combine(day, t_start)),
                    end=common.to_iso(datetime.combine(day, t_end))))
            day += timedelta(days=1)
    return lessons


# ---------- .ics ----------
def _as_msk(value) -> datetime | None:
    """Приводит DTSTART/DTEND к Москве. Целодневные события (date) пропускаем."""
    if isinstance(value, datetime):
        return value.replace(tzinfo=common.MSK) if value.tzinfo is None \
            else value.astimezone(common.MSK)
    return None


def _expand_ics(path: str, errors: list[str]) -> list[Lesson]:
    import icalendar
    import recurring_ical_events

    cal = icalendar.Calendar.from_ical(Path(path).read_bytes())
    start, end, _, holidays = _semester()
    window_from = datetime.combine(start, time(0, 0), tzinfo=common.MSK)
    window_to = datetime.combine(end + timedelta(days=1), time(0, 0), tzinfo=common.MSK)
    lessons: list[Lesson] = []
    events = recurring_ical_events.of(cal).between(window_from, window_to)
    for n, ev in enumerate(sorted(events, key=lambda e: str(e.get("DTSTART").dt)), start=1):
        summary = str(ev.get("SUMMARY", "")).strip()
        begin, finish = _as_msk(ev.get("DTSTART").dt), None
        if ev.get("DTEND") is not None:
            finish = _as_msk(ev.get("DTEND").dt)
        elif ev.get("DURATION") is not None:
            finish = begin + ev.get("DURATION").dt if begin else None
        if begin is None:
            continue                                            # целодневное событие, не пара
        if finish is None:
            errors.append(f"Событие «{summary}» {begin:%d.%m.%Y}: нет времени окончания")
            continue
        disc = config.find_discipline(summary + " " + str(ev.get("DESCRIPTION", "")))
        if disc is None:
            errors.append(f'Событие «{summary}» {begin:%d.%m.%Y}: не найдена дисциплина, '
                          f"добавьте алиас в disciplines.yaml")
            continue
        if begin.date() in holidays:
            continue
        teacher = None
        m = re.search(r"(?:преподаватель|препод\.?)\s*[:\-]\s*(.+)", str(ev.get("DESCRIPTION", "")),
                      re.I)
        if m:
            teacher = m.group(1).splitlines()[0].strip() or None
        lessons.append(Lesson(
            discipline=disc, lesson_type=_lesson_type(summary), teacher=teacher,
            room=_clean(ev.get("LOCATION")), start=common.to_iso(begin), end=common.to_iso(finish)))
    return lessons


# ---------- публичная функция ----------
def load_schedule(path: str, dry_run: bool = False) -> ScheduleImport:
    """Читает .ics или .xlsx и заменяет расписание в базе.

    dry_run=True ничего не записывает: предпросмотр занятий и список ошибок.
    При любой ошибке запись не выполняется. ok=True значит «ошибок нет» (при dry_run
    запись возможна, без dry_run она выполнена). Исключения наружу не выходят.
    """
    errors: list[str] = []
    lessons: list[Lesson] = []
    p = Path(path)
    ext = p.suffix.lower()
    if not p.exists():
        errors.append(f"Файл не найден: {path}")
    elif ext not in (".ics", ".xlsx", ".xlsm"):
        errors.append(f"Формат «{ext or 'без расширения'}» не поддерживается: "
                      f"загрузите .ics или .xlsx")
    else:
        try:
            lessons = _expand_ics(path, errors) if ext == ".ics" else _expand_excel(path, errors)
        except Exception as e:                                  # битый файл, чужой формат
            errors.append(f"Не удалось прочитать файл «{p.name}»: {e}")
            lessons = []
    if not errors:
        _check(lessons, errors)
    if not errors and not lessons:
        errors.append("В файле не найдено ни одного занятия в границах семестра")
    if errors:
        return ScheduleImport(ok=False, lessons=lessons, errors=errors, replaced=0)
    lessons.sort(key=lambda x: (x.start, x.discipline))
    replaced = db.count_lessons() if dry_run else db.replace_lessons(lessons)
    return ScheduleImport(ok=True, lessons=lessons, errors=[], replaced=replaced)


def make_template(path: str) -> None:
    """Создаёт шаблон Excel-расписания (docs/schedule_template.xlsx) с примерами строк."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Расписание"
    ws.append(REQUIRED_COLUMNS)
    ws.append(["Электротехника", "лекция", "Иванов И. И.", "А-201", "вторник", "08:30", "10:05", "каждая"])
    ws.append(["Электротехника", "лабораторная", "Петров П. П.", "Б-105", "четверг", "10:15", "11:50", "нечёт"])
    ws.append(["Модели данных", "семинар", "Сидорова А. А.", "В-310", "пятница", "12:00", "13:35", "чёт"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCDEFGH", (22, 16, 20, 12, 14, 9, 9, 10)):
        ws.column_dimensions[col].width = width
    note = wb.create_sheet("Подсказка")
    note.append(["Колонка", "Допустимые значения"])
    note.append(["дисциплина", "название или алиас из config/disciplines.yaml"])
    note.append(["тип", "лекция, семинар, лабораторная или другое"])
    note.append(["день_недели", "пн, вт, ср, чт, пт, сб, вс или полное название"])
    note.append(["начало, конец", "время в формате ЧЧ:ММ"])
    note.append(["неделя", "чёт, нечёт или каждая (первая неделя семестра задаётся в config.yaml)"])
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


if __name__ == "__main__":                                      # python -m src.tracker.schedule_import файл
    import json
    import sys
    if len(sys.argv) > 1 and sys.argv[1] == "--template":
        make_template(sys.argv[2] if len(sys.argv) > 2 else "docs/schedule_template.xlsx")
    else:
        dry = "--dry-run" in sys.argv
        files = [a for a in sys.argv[1:] if not a.startswith("--")]
        res = load_schedule(files[0], dry_run=dry)
        print(json.dumps({"ok": res.ok, "errors": res.errors, "replaced": res.replaced,
                          "lessons": len(res.lessons)}, ensure_ascii=False, indent=2))
